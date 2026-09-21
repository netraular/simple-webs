#!/usr/bin/env python3
"""
Construeix compra.json i compra-bcn-barris.json a partir de dades oficials.

No estima, no interpola i no converteix cap preu: si una font no publica el
valor d'un municipi o barri, aquell municipi/barri no s'inclou al resultat.

L'univers de municipis es el de `municipis.json` d'aquest mateix directori
(92 municipis amb el cap de municipi a <= 32 km de la Placa de Catalunya), de
manera que `codi_ine`, `nom` i `dist_bcn_km` son exactament els mateixos que
a la resta de fitxers del projecte.

Fonts (vegeu compra.SOURCES.md):
  A) Generalitat de Catalunya, Secretaria d'Habitatge - "Compravendes
     d'habitatge registrades" (origen: Col.legi de Registradors).
     Preu per m2 construit d'habitatge USAT. -> metrica principal.
  B) Ministerio de Transportes / Vivienda - "Valor tasado medio de vivienda
     libre" de municipis > 25.000 hab. -> nomes com a creuament.

Fitxers d'entrada esperats a $WORK (per defecte /tmp), tal com es descarreguen
de les URLs documentades a compra.SOURCES.md:
  g_MUN_acum1any_2025.xlsx       (font A, municipis de Catalunya)
  f_BCN_usat_acum1any_2025.xlsx  (font A, districtes i barris de Barcelona)
  vt_mun.xls                     (font B)
"""

import json
import os
import re
import unicodedata

import openpyxl
import xlrd

WORK = os.environ.get("WORK", "/tmp")
OUT = os.path.dirname(os.path.abspath(__file__))

CODI_BCN = "08019"
SHEET_MUN = "4t25acum_1any"  # acumulat de 4 trimestres tancat al 4t de 2025
SHEET_VT = "T2A2026"  # darrer trimestre publicat pel Ministerio

URL_MUN = (
    "https://habitatge.gencat.cat/web/.content/home/dades/estadistiques/"
    "01_Estadistiques_de_construccio_i_mercat_immobiliari/02_Compravenda_i_preu_de_venda/"
    "02_Compravendes_d_habitatges_registrades_i_el_preu_de_venda/2025/MUN_acum1any_2025.xlsx"
)
URL_BCN = (
    "https://habitatge.gencat.cat/web/.content/home/dades/estadistiques/"
    "01_Estadistiques_de_construccio_i_mercat_immobiliari/02_Compravenda_i_preu_de_venda/"
    "02_Compravendes_d_habitatges_registrades_i_el_preu_de_venda/2025/BCN_usat_acum1any_2025.xlsx"
)
URL_VT = "https://apps.fomento.gob.es/BoletinOnline2/sedal/35103500.XLS"

FONT_A = (
    "Generalitat de Catalunya, Secretaria d'Habitatge i Inclusió Social — "
    "Estadística de compravendes d'habitatge registrades "
    "(origen de les dades: Col·legi de Registradors de la Propietat)"
)
FONT_B = (
    "Ministerio de Transportes y Movilidad Sostenible / Ministerio de Vivienda y "
    "Agenda Urbana — Boletín Estadístico Online, taula 35103500 «Valor tasado de "
    "vivienda libre de los municipios mayores de 25.000 habitantes»"
)

METRICA_A = (
    "preu mitjà de venda per m² construït d'habitatge USAT (segona mà), "
    "calculat sobre les compravendes efectivament registrades"
)
METRICA_B = (
    "valor mitjà de taxació d'habitatge lliure (Ordre ECO/805/2003); és una "
    "valoració pericial, NO un preu de transacció"
)

# Equivalències ortogràfiques entre la grafia castellanitzada del Ministerio i
# el nom oficial. Només resolen com s'escriu el nom; no aporten cap dada nova.
ALIES_VT = {"santa coloma gramanet": "08245"}


def norm(s):
    """Normalitza un nom de municipi per comparar-lo entre fonts."""
    s = str(s or "")
    s = re.sub(r"\s*\(\d+\)\s*$", "", s)  # marca de nota al peu: "Cervelló (2)"
    s = re.sub(r"\s*\((?:el|la|els|les|l['’]|es|sa)\)\s*$", "", s, flags=re.I)
    s = re.sub(r",\s*(?:el|la|els|les|l['’]|es|sa)\s*$", "", s, flags=re.I)
    s = re.sub(r"^(?:el|la|els|les|l['’]|es|sa)\s+", "", s, flags=re.I)
    s = re.sub(r"^l['’]", "", s, flags=re.I)
    s = unicodedata.normalize("NFKD", s)
    s = "".join(c for c in s if not unicodedata.combining(c)).lower()
    s = re.sub(r"[^a-z0-9 ]+", " ", s)
    return re.sub(r"\s+", " ", s).strip()


def num(v):
    """float si el valor és un número estrictament positiu; si no, None.

    A l'origen, «n.d.», «n.r» i 0 signifiquen «sense dada publicada».
    """
    if v is None:
        return None
    if isinstance(v, str):
        v = v.strip().replace(",", ".")
        if not re.fullmatch(r"-?\d+(\.\d+)?", v):
            return None
    try:
        v = float(v)
    except (TypeError, ValueError):
        return None
    return v if v > 0 else None


def ival(v):
    n = num(v)
    return int(round(n)) if n is not None else None


def r1(v):
    return round(v, 1) if v is not None else None


# --------------------------------------------------------------------------
# Univers de municipis: el canònic del projecte
# --------------------------------------------------------------------------
canonics = json.load(open(os.path.join(OUT, "municipis.json"), encoding="utf-8"))

# --------------------------------------------------------------------------
# Font B — valor tasat del Ministerio (creuament)
# --------------------------------------------------------------------------
vt_sheet = xlrd.open_workbook(os.path.join(WORK, "vt_mun.xls")).sheet_by_name(SHEET_VT)
periode_vt = re.sub(r"\(\*\)", "", str(vt_sheet.cell_value(11, 1))).strip()
vt_by_norm, prov = {}, None
for r in range(17, vt_sheet.nrows):
    row = [c.value for c in vt_sheet.row(r)]
    p, m = str(row[1]).strip(), str(row[2]).strip()
    if p:
        prov = p
    if not m or prov != "Barcelona":
        continue
    vt_by_norm[norm(m)] = {
        "eur_m2_mes_de_5_anys_antiguitat": num(row[4]),
        "eur_m2_total": num(row[5]),
        "nombre_tasacions_mes_de_5_anys": ival(row[8]),
        "nombre_tasacions_total": ival(row[9]),
        "periode": periode_vt,
        "metrica": METRICA_B,
        "font": FONT_B,
        "font_url": URL_VT,
    }

vt_by_ine = {}
for m in canonics:
    d = vt_by_norm.get(norm(m["nom"]))
    if d:
        vt_by_ine[m["codi_ine"]] = d
for alias, ine in ALIES_VT.items():
    if alias in vt_by_norm:
        vt_by_ine.setdefault(ine, vt_by_norm[alias])

# --------------------------------------------------------------------------
# Font A — compravendes registrades per municipi
# --------------------------------------------------------------------------
ws = openpyxl.load_workbook(
    os.path.join(WORK, "g_MUN_acum1any_2025.xlsx"), data_only=True
)[SHEET_MUN]
cobertura_mun = str(ws.cell(1, 1).value or "").strip()
periode_mun = re.sub(r"^Per[ií]ode:\s*", "", str(ws.cell(2, 1).value or "").strip())

# Notes al peu del full: "(1) ...", "(2) ...". Poden estar a la columna 1 o a la 2.
notes_peu = {}
for r in range(6, ws.max_row + 1):
    for col in (1, 2):
        mnote = re.match(r"^\((\d+)\)\s*(.+)$", str(ws.cell(r, col).value or "").strip())
        if mnote:
            notes_peu[mnote.group(1)] = mnote.group(2).strip()

# Files de dades indexades per codi INE
files = {}
for r in range(6, ws.max_row + 1):
    codi = ws.cell(r, 1).value
    if not codi or not str(codi).strip().isdigit():
        continue
    nom_font = str(ws.cell(r, 2).value or "").strip()
    marca = re.search(r"\((\d+)\)\s*$", nom_font)
    files[str(codi).strip().zfill(5)] = {
        "eur_m2_usat": num(ws.cell(r, 18).value),
        "eur_m2_nou": num(ws.cell(r, 17).value),
        "eur_m2_total": num(ws.cell(r, 19).value),
        "n_usat": ival(ws.cell(r, 5).value),
        "n_total": ival(ws.cell(r, 6).value),
        "sup_usat": num(ws.cell(r, 10).value),
        "preu_total_usat_milers": num(ws.cell(r, 14).value),
        "nota": notes_peu.get(marca.group(1)) if marca else None,
    }

municipis, sense_dada = [], []
for m in canonics:
    ine = m["codi_ine"]
    d = files.get(ine)
    if d is None or d["eur_m2_usat"] is None:
        sense_dada.append(
            {
                "codi_ine": ine,
                "nom": m["nom"],
                "motiu": (
                    "no apareix a la font (municipi de menys de 2.000 habitants "
                    "o integrat en un altre municipi)"
                    if d is None
                    else "la font no publica preu per a aquest municipi"
                ),
            }
        )
        continue

    rec = {
        "codi_ine": ine,
        "nom": m["nom"],
        "compra_eur_m2": r1(d["eur_m2_usat"]),
        "metrica": METRICA_A,
        "periode": periode_mun,
        "any": 2025,
        "font": FONT_A,
        "font_url": URL_MUN,
        "nombre_operacions": d["n_usat"],
        "superficie_mitjana_m2": d["sup_usat"],
        "compra_eur_total": (
            round(d["preu_total_usat_milers"] * 1000, 2)
            if d["preu_total_usat_milers"]
            else None
        ),
        "compra_eur_m2_obra_nova": r1(d["eur_m2_nou"]),
        "compra_eur_m2_tots_els_habitatges": r1(d["eur_m2_total"]),
        "nombre_operacions_total": d["n_total"],
        "dist_bcn_km": m["dist_bcn_km"],
        "comarca": m.get("comarca"),
        "valor_tasat_ministeri": vt_by_ine.get(ine),
    }
    if d["nota"]:
        rec["nota"] = d["nota"]
    municipis.append(rec)

compra = {
    "meta": {
        "descripcio": (
            "Preu de compra d'habitatge en euros per m² construït als municipis de "
            "l'entorn de Barcelona. L'univers de municipis és el de municipis.json "
            "(cap de municipi a 32 km o menys de la plaça de Catalunya)."
        ),
        "camp_principal": "compra_eur_m2",
        "metrica_principal": METRICA_A,
        "periode": periode_mun,
        "any": 2025,
        "font_principal": FONT_A,
        "font_principal_url": URL_MUN,
        "cobertura_font_principal": cobertura_mun,
        "creuament": {
            "camp": "valor_tasat_ministeri",
            "metrica": METRICA_B,
            "periode": periode_vt,
            "font": FONT_B,
            "font_url": URL_VT,
            "cobertura": "només municipis de més de 25.000 habitants",
            "nombre_municipis_amb_creuament": sum(
                1 for x in municipis if x["valor_tasat_ministeri"]
            ),
        },
        "nombre_municipis": len(municipis),
        "municipis_de_referencia": len(canonics),
        "sense_dada": sense_dada,
        "avis": (
            "Cap valor no està estimat, interpolat ni convertit. Els municipis sense "
            "preu publicat es llisten a meta.sense_dada i no apareixen a la llista "
            "«municipis». compra_eur_m2 (compravendes registrades) i "
            "valor_tasat_ministeri (taxacions) mesuren coses diferents i no són "
            "intercanviables."
        ),
    },
    "municipis": municipis,
}
with open(os.path.join(OUT, "compra.json"), "w", encoding="utf-8") as f:
    json.dump(compra, f, ensure_ascii=False, indent=2)
    f.write("\n")

# --------------------------------------------------------------------------
# Font A-bis — districtes i barris de Barcelona ciutat
# --------------------------------------------------------------------------
wb = openpyxl.load_workbook(
    os.path.join(WORK, "f_BCN_usat_acum1any_2025.xlsx"), data_only=True
)[SHEET_MUN]
periode_bcn = re.sub(r"^Per[ií]ode:\s*", "", str(wb.cell(2, 1).value or "").strip())

ciutat, districtes, barris, seccio, exclosos = None, [], [], None, []
for r in range(5, wb.max_row + 1):
    nom = str(wb.cell(r, 2).value or "").strip()
    if not nom or nom.startswith("*"):
        continue
    if nom.startswith("Districtes municipals"):
        seccio = "districte"
        continue
    if nom.startswith("Barris de Barcelona"):
        seccio = "barri"
        continue

    eur_m2 = num(wb.cell(r, 6).value)
    preu_total = num(wb.cell(r, 5).value)
    rec = {
        "nom": nom.rstrip("*").strip(),
        "compra_eur_m2": r1(eur_m2),
        "nombre_operacions": ival(wb.cell(r, 3).value),
        "superficie_mitjana_m2": num(wb.cell(r, 4).value),
        "compra_eur_total": round(preu_total * 1000, 2) if preu_total else None,
        "compra_eur_m2_maxim": num(wb.cell(r, 7).value),
        "compra_eur_m2_minim": num(wb.cell(r, 8).value),
    }
    if seccio is None:
        ciutat = {"codi_ine": CODI_BCN, **rec}
        continue
    if eur_m2 is None:
        exclosos.append({"tipus": seccio, "nom": rec["nom"]})
        continue
    codi = wb.cell(r, 1).value
    rec = {"codi": str(codi).strip() if codi is not None else None, **rec}
    (districtes if seccio == "districte" else barris).append(rec)

bcn = {
    "meta": {
        "descripcio": (
            "Preu de compra d'habitatge usat (segona mà) per districte i per barri "
            "de la ciutat de Barcelona."
        ),
        "camp_principal": "compra_eur_m2",
        "metrica_principal": METRICA_A,
        "periode": periode_bcn,
        "any": 2025,
        "font": FONT_A,
        "font_url": URL_BCN,
        "codis": (
            "«codi» és el codi oficial de districte (1-10) o de barri (1-73) de "
            "l'Ajuntament de Barcelona; NO és un codi INE."
        ),
        "nombre_districtes": len(districtes),
        "nombre_barris": len(barris),
        "barris_totals_a_la_ciutat": 73,
        "exclosos_sense_dada": exclosos,
        "avis": (
            "Cap valor no està estimat. Els barris on la font publica «n.d.» o 0 "
            "(mostra insuficient) s'han exclòs i es llisten a exclosos_sense_dada. "
            "El total de la ciutat no és la suma dels districtes (vegeu la nota de "
            "la font)."
        ),
    },
    "ciutat": ciutat,
    "districtes": districtes,
    "barris": barris,
}
with open(os.path.join(OUT, "compra-bcn-barris.json"), "w", encoding="utf-8") as f:
    json.dump(bcn, f, ensure_ascii=False, indent=2)
    f.write("\n")

# --------------------------------------------------------------------------
# Informe
# --------------------------------------------------------------------------
print(f"compra.json            : {len(municipis)}/{len(canonics)} municipis amb dada")
print(f"  període              : {periode_mun}")
print(f"  creuament Ministerio : {sum(1 for x in municipis if x['valor_tasat_ministeri'])}")
print(f"  sense dada           : {[(x['nom']) for x in sense_dada]}")
print(f"compra-bcn-barris.json : {len(districtes)} districtes, {len(barris)}/73 barris")
print(f"  exclosos             : {[x['nom'] for x in exclosos]}")
print("\nComprovació de cordura — compra_eur_m2 (habitatge usat, 2025):")
for n in [
    "Barcelona",
    "Sant Cugat del Vallès",
    "Terrassa",
    "Mataró",
    "Castelldefels",
    "Sabadell",
]:
    for x in municipis:
        if x["nom"] == n:
            vt = x["valor_tasat_ministeri"]
            v = vt["eur_m2_mes_de_5_anys_antiguitat"] if vt else "—"
            print(
                f"  {n:24s} {x['compra_eur_m2']:>7.1f} €/m²  n={x['nombre_operacions']:>5}  "
                f"{x['dist_bcn_km']:>5} km   [valor tasat Min.: {v}]"
            )
